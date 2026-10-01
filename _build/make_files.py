#!/usr/bin/env python3
"""Генерирует скачиваемые файлы в site/files/."""
import pathlib, subprocess, zipfile, datetime, tempfile, csv
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

import os
FILES = pathlib.Path(os.environ.get("OUT_DIR", "files")); FILES.mkdir(parents=True, exist_ok=True)

# 1) PDF чек-лист — через headless Chrome
check_html = """<!doctype html><html lang="ru"><meta charset="utf-8"><style>
body{font-family:'DejaVu Sans',Arial,sans-serif;margin:18mm 16mm;color:#1d2433;font-size:11pt}
h1{color:#3b5bdb;margin:0 0 4px;font-size:20pt}.sub{color:#5b6478;margin:0 0 14px}
h2{font-size:13pt;margin:14px 0 6px;border-bottom:2px solid #dbe4ff;padding-bottom:3px}
li{list-style:none;margin:5px 0}li::before{content:"☐  ";color:#3b5bdb}
.f{margin-top:18px;color:#868e96;font-size:9pt}</style>
<h1>Чек-лист подготовки к сессии</h1><p class="sub">Студент.Навигатор — полезное для студентов</p>
<h2>За 4 недели</h2><ul><li>Выписать все экзамены, зачёты, курсовые и даты</li><li>Узнать условия допуска по каждой дисциплине</li><li>Составить список «хвостов» и план их закрытия</li></ul>
<h2>За 3 недели</h2><ul><li>Собрать вопросы к экзаменам и конспекты</li><li>Сделать карту тем: зелёный / жёлтый / красный</li><li>Согласовать с руководителем черновик курсовой</li></ul>
<h2>За 2 недели</h2><ul><li>Каждый день 4–6 «помидоров» по плану</li><li>Повторять пройденное через 1, 3 и 7 дней</li><li>Сдать курсовые и отчёты, получить допуск</li></ul>
<h2>За неделю</h2><ul><li>Пробный экзамен: случайный билет, 20 минут, ответ вслух</li><li>Отметить «плавающие» темы и вернуться к ним</li></ul>
<h2>Накануне</h2><ul><li>Лёгкое повторение определений и формул</li><li>Подготовить зачётку, студенческий, ручки</li><li>Лечь спать вовремя: 7–9 часов сна</li></ul>
<h2>На экзамене</h2><ul><li>Прочитать все вопросы, начать с самого понятного</li><li>План ответа: определение → суть → пример → связи</li></ul>
<p class="f">Полная версия гида — на странице «Гид по сессии». © Студент.Навигатор</p></html>"""
with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
    f.write(check_html); tmp = f.name
subprocess.run(["google-chrome", "--headless=new", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                f"--print-to-pdf={FILES/'chek-list-sessii.pdf'}", f"file://{tmp}"], check=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

# 2) DOCX шаблон отчёта
d = Document()
sec = d.sections[0]
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(14)
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.first_line_indent = Mm(12.5)
for name in ("Heading 1", "Heading 2"):
    h = d.styles[name]; h.font.name = "Times New Roman"; h.font.size = Pt(14 if name == "Heading 2" else 16); h.font.bold = True
    h.font.color.rgb = None
def c(text, bold=False, size=None):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = 0
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    return p
for t in ["Министерство науки и высшего образования Российской Федерации", "Федеральное государственное автономное образовательное учреждение высшего образования", "«МОСКОВСКИЙ ПОЛИТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ»", "(МОСКОВСКИЙ ПОЛИТЕХ)", "", "Факультет ________", "Кафедра «________»", "", "", ""]:
    c(t, bold=t.startswith("«") or t.startswith("(М"))
c("Лабораторная работа № __", bold=True, size=16); c("Тема: «________»"); c("по дисциплине «________»")
for _ in range(4): c("")
p = d.add_paragraph("Студент ________ (личная подпись) (И. О. Фамилия)"); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p = d.add_paragraph("Преподаватель ________ (И. О. Фамилия)"); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for _ in range(3): c("")
c(f"Москва {datetime.date.today().year}")
d.add_page_break()
c("СОДЕРЖАНИЕ", bold=True)
d.add_paragraph("(Вставьте автоматическое оглавление: Ссылки → Оглавление)")
d.add_page_break()
c("ВВЕДЕНИЕ", bold=True)
d.add_paragraph("Актуальность, цель и задачи работы. Объект и предмет исследования.")
d.add_page_break()
d.add_heading("1 Название первого раздела", level=1)
d.add_heading("1.1 Название подраздела", level=2)
d.add_paragraph("Текст раздела. Ссылка на рисунок в тексте обязательна (рисунок 1).")
p = d.add_paragraph("[Место для рисунка]"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = d.add_paragraph("Рисунок 1 — Название рисунка"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = 0
p = d.add_paragraph("Таблица 1 — Название таблицы"); p.paragraph_format.first_line_indent = 0
tb = d.add_table(rows=3, cols=3); tb.style = "Table Grid"
for i, h in enumerate(["Показатель", "Значение", "Комментарий"]): tb.rows[0].cells[i].text = h
d.add_page_break()
c("ЗАКЛЮЧЕНИЕ", bold=True); d.add_paragraph("Основные результаты и выводы.")
d.add_page_break()
c("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", bold=True)
d.add_paragraph("1. Фамилия, И. О. Название книги / И. О. Фамилия. — Москва : Издательство, 2024. — 200 с. — ISBN ...")
d.add_paragraph("2. Название страницы : [сайт]. — URL: https://example.ru (дата обращения: 01.10.2026). — Текст : электронный.")
d.core_properties.author = "Студент.Навигатор"; d.core_properties.title = "Шаблон отчёта по ГОСТ 7.32-2017"
d.save(FILES / "shablon-otcheta-gost-7.32.docx")

# 3) ICS календарь
y = datetime.date.today().year
events = [(f"{y}0901", "Начало осеннего семестра"), (f"{y}1020", "Контрольная неделя (ориентир)"), (f"{y}1201", "Сдать курсовые работы (ориентир)"),
          (f"{y}1215", "Зачётная неделя (ориентир)"), (f"{y+1}0109", "Начало зимней сессии (ориентир)"), (f"{y+1}0207", "Начало весеннего семестра (ориентир)")]
now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Student Navigator//RU", "CALSCALE:GREGORIAN", "X-WR-CALNAME:Семестр — Студент.Навигатор"]
for i, (dt, s) in enumerate(events):
    lines += ["BEGIN:VEVENT", f"UID:sn-{dt}-{i}@student-navigator", f"DTSTAMP:{now}", f"DTSTART;VALUE=DATE:{dt}", f"SUMMARY:{s}",
              "DESCRIPTION:Даты ориентировочные — сверьте с графиком учебного процесса вашего вуза.", "BEGIN:VALARM", "TRIGGER:-P3D", "ACTION:DISPLAY", f"DESCRIPTION:{s}", "END:VALARM", "END:VEVENT"]
lines.append("END:VCALENDAR")
(FILES / "kalendar-semestra.ics").write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")

# 4) CSV планер (UTF-8 с BOM — чтобы Excel открыл кириллицу)
with open(FILES / "planer-sessii.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(["Дисциплина", "Форма контроля", "Дата", "Допуск получен (да/нет)", "Хвосты", "Тем всего", "Тем «зелёных»", "Готовность, %", "Комментарий"])
    w.writerow(["Пример: Веб-аналитика", "Экзамен", "15.01", "да", "—", "30", "18", "=G2/F2*100", "повторить Вебвизор и цели"])
    for _ in range(12): w.writerow([""] * 9)

# 5) ZIP со всеми файлами
with zipfile.ZipFile(FILES / "student-navigator-pack.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for fn in ["chek-list-sessii.pdf", "shablon-otcheta-gost-7.32.docx", "kalendar-semestra.ics", "planer-sessii.csv"]:
        z.write(FILES / fn, fn)
for p in sorted(FILES.iterdir()): print(p.name, p.stat().st_size)
